"""Audit evidence, write checksums, and create the source/document release."""
from pathlib import Path
import csv
import hashlib
import json
import platform
import re
import subprocess
import zipfile
import numpy as np
from PIL import Image,ImageStat
from pypdf import PdfReader
from docx import Document
ROOT=Path(__file__).resolve().parents[1]


def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def included(path):
    relative=path.relative_to(ROOT)
    return (path.is_file() and not any(x in {'.venv','__pycache__','mplconfig','runtime_cache'} for x in relative.parts)
            and path.suffix not in {'.so','.dll','.pyc'} and path.name!='SHA256SUMS.txt')


def main():
    d=json.loads((ROOT/'results/validation.json').read_text())
    assert d['summary']=={'total':61,'passed':61,'failed':0}
    assert 'PASS:' in (ROOT/'results/viewer_smoke.log').read_text()
    assert "'failed': 0" in (ROOT/'results/setup_host_repeat.log').read_text()
    assert 'Plot failed!' not in (ROOT/'results/octave_run.log').read_text()
    octave=ROOT/'matlab/octave_output'
    assert 'GNU Octave 9.4.0' in (octave/'RUN_COMPLETED.txt').read_text()
    plots=list(octave.glob('*.png'));assert len(plots)==30
    for p in plots:assert ImageStat.Stat(Image.open(p).convert('L')).stddev[0]>10
    a=np.loadtxt(ROOT/'results/fft_c_and_dif.csv',delimiter=',',skiprows=1)
    b=np.loadtxt(octave/'fft_reference.csv',delimiter=',')
    assert np.array_equal(a[:,4:6],b[:,3:5]) and np.array_equal(a[:,6:8],b[:,5:7])
    with (ROOT/'results/filter_coefficients.csv').open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
    assert len(rows)==28*257
    q=np.array([int(r['q15_integer']) for r in rows]).reshape(28,257)
    assert np.array_equal(q,np.loadtxt(octave/'coefficients_q15.csv',delimiter=','))
    original=Path('/workspace/attachments/fba9db3d-3da7-44fc-b80d-e915c3b35231/v61.zip')
    original_sha=digest(ROOT/'original_v61.zip')
    if original.exists():assert digest(original)==original_sha
    for name in ['课程设计报告.docx','操作手册.docx']:
        doc=Document(ROOT/'docs'/name);assert len(doc.paragraphs)>80
        with zipfile.ZipFile(ROOT/'docs'/name) as z:assert z.testzip() is None
    assert len(Document(ROOT/'docs/课程设计报告.docx').inline_shapes)==4
    words=json.loads((ROOT/'results/report_word_count.json').read_text())
    assert 5000<=words['main_body_chinese_characters']<=10000
    for p in ROOT.rglob('*'):
        if p.is_file() and p.suffix.lower() in {'.bat','.pjt'}:
            raw=p.read_bytes().replace(b'\r\n',b'\n')
            p.write_bytes(raw.replace(b'\n',b'\r\n'))
    evidence={
        'native_C_algorithm_checks':d['summary'],
        'PC_GUI_callbacks':{'keys_exercised':9,'snapshot_files_refreshed':2,'backend':'headless Agg'},
        'Octave':{'version':'9.4.0','completed':True,'plots':30,'coefficients_equal_to_C_tables':True,'DIT_and_DIF_equal_to_C_outputs':True},
        'filters':{'banks':28,'taps_each':257,'fs_Hz':32000,'worst_stop_attenuation_dB':min(x['stop_attenuation_dB'] for x in d['filters']),'max_pass_ripple_dB':max(x['ripple_dB'] for x in d['filters'])},
        'report':{**words,'PDF_pages':len(PdfReader(ROOT/'docs/课程设计报告.pdf').pages)},
        'coefficient_appendix_pages':len(PdfReader(ROOT/'docs/FIR滤波系数表.pdf').pages),
        'original_archive_sha256':original_sha,
        'not_run':['TI C55x compilation/assembly/link','Windows host GUI/build','MATLAB commercial runtime','physical JTAG/keypad/board','FLASH programming or power-cycle test'],
        'assembly_kind':'Independent DIF description -> cl55-generated C55x assembly -> reassembly; not handwritten. Target assembly artifact not yet generated.',
        'hardware':'Source identifies ICETEK C5509; actual board unconfirmed; large manual named DSP28335 not read.',
        'native_platform':platform.platform(),
    }
    (ROOT/'results/final_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    text=['# 验证结果说明','',f"便携 C 最终 {d['summary']['total']} 项检查全部通过；安装脚本在新虚拟环境及重复执行时均成功。",'',
          '界面使用无界面渲染实际执行了全部 9 个按键回调，并刷新两个 CCS 格式示例快照。此记录不是 Windows 实测，也不是板卡连接测试。','',
          'Matlab 脚本由 GNU Octave 9.4.0 实际执行，产生 30 张非空图及数值文件。全量系数与 C 表一致，两种 FFT 模型与对应 C 输出逐项一致。日志中的 gnuplot 维护状态提醒不影响此次批量绘图；最终运行退出码为 0，图形内容已检查。','',
          '以下数值来自实际便携 C 输出，频谱采用 FFT/N：','',
          '| 模式 | 8 kHz 载波（整数计数） | 7 kHz 边带 | 9 kHz 边带 |','|---|---:|---:|---:|']
    for x in d['modulation']:text.append(f"| {x['mode']} | {x['carrier_count']:.5f} | {x['lower_count']:.5f} | {x['upper_count']:.5f} |")
    text+=['','| 滤波器 | 通带峰峰波动 dB | 最差阻带衰减 dB |','|---|---:|---:|']
    for x in d['filters']:text.append(f"| {x['name']} | {x['ripple_dB']:.6f} | {x['stop_attenuation_dB']:.6f} |")
    text+=['','TI 目标编译/汇编链接、MATLAB 本体、真实板卡按键/下载和 FLASH 均未验证。汇编采用编译器生成路径，不是手写汇编；若教师要求手写，这一要求仍需补充。',
           '',f"论文正文汉字统计：{words['main_body_chinese_characters']}（不含摘要、文献及附录）；Word 为可编辑版本，PDF 为阅读版。"]
    (ROOT/'docs/仿真与验证说明.md').write_text('\n'.join(text)+'\n',encoding='utf-8')
    files=sorted(p for p in ROOT.rglob('*') if included(p))
    manifest='\n'.join(f'{digest(p)}  {p.relative_to(ROOT).as_posix()}' for p in files)+'\n'
    (ROOT/'SHA256SUMS.txt').write_text(manifest,encoding='utf-8')
    out=ROOT.parent/'数字系统综合设计_代码与文档.zip'
    with zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        for p in files+[ROOT/'SHA256SUMS.txt']:z.write(p,'course_design/'+p.relative_to(ROOT).as_posix())
    with zipfile.ZipFile(out) as z:
        assert z.testzip() is None
        assert not any('/.venv/' in x or x.endswith('.so') for x in z.namelist())
    out.with_suffix('.zip.sha256').write_text(digest(out)+'  '+out.name+'\n',encoding='utf-8')
    print(json.dumps({'archive':str(out),'bytes':out.stat().st_size,'files':len(files)+1,'report_CJK_characters':words['main_body_chinese_characters']},ensure_ascii=False))


if __name__=='__main__':main()
