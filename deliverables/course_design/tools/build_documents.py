"""Rebuild reviewable DOCX/PDF documents (pandoc, python-docx, reportlab)."""
from pathlib import Path
import re
import subprocess
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Image,PageBreak,KeepTogether
from PIL import Image as PILImage
ROOT=Path(__file__).resolve().parents[1]


def font(style,zh='宋体',size=12,bold=False):
    style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold
    rpr=style.element.get_or_add_rPr();rfonts=rpr.find(qn('w:rFonts'))
    if rfonts is None:rfonts=OxmlElement('w:rFonts');rpr.insert(0,rfonts)
    rfonts.set(qn('w:eastAsia'),zh)


def docx_build(md):
    out=md.with_suffix('.docx')
    subprocess.run(['pandoc',str(md),'-o',str(out),'--resource-path='+str(md.parent)],check=True)
    d=Document(out)
    # Pandoc capitalizes built-in style names; python-docx expects Word's
    # canonical lowercase names for heading lookup.
    for st in d.styles:
        if re.fullmatch(r'Heading [1-9]',st.name):
            st.element.find(qn('w:name')).set(qn('w:val'),st.name.lower())
    for p in d.paragraphs:
        if p.style.name=='Heading 2':p.style=d.styles['Heading 1']
        elif p.style.name=='Heading 3':p.style=d.styles['Heading 2']
    d.paragraphs[0].style=d.styles['Title']
    for sec in d.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=Cm(2.5);sec.bottom_margin=Cm(2.5)
        sec.left_margin=Cm(2.7);sec.right_margin=Cm(2.5)
        p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
    font(d.styles['Normal']);d.styles['Normal'].paragraph_format.line_spacing=1.5
    d.styles['Normal'].paragraph_format.space_after=Pt(4)
    for k,size in [(1,16),(2,14),(3,12)]:
        font(d.styles[f'Heading {k}'],'黑体',size,True)
        d.styles[f'Heading {k}'].paragraph_format.keep_with_next=True
    for p in d.paragraphs:
        if p.style.name in ('Normal','Body Text','First Paragraph') and p.text.strip():p.paragraph_format.first_line_indent=Pt(24)
        if p.text.startswith('图 ') or p.text.startswith('图1'):
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=0
        if p.style.name in ('Source Code','Verbatim Char','Code'):
            for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
    if '课程设计报告' in md.stem:
        first=d.paragraphs[0];first.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in first.runs:r.font.size=Pt(20)
        # Keep a clean cover, then abstract pages; editable Word TOC.
        abstract=next(p for p in d.paragraphs if p.text=='摘要')
        abstract.paragraph_format.page_break_before=True
        en=next(p for p in d.paragraphs if p.text=='Abstract')
        en.paragraph_format.page_break_before=True
        chapter=next(p for p in d.paragraphs if p.text.startswith('1　研究'))
        toc=chapter.insert_paragraph_before('目　录','Heading 1')
        toc.paragraph_format.page_break_before=True
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
        hint=chapter.insert_paragraph_before('在 Word 中右键目录选择“更新域”，生成本机页码。')
        hint._p.append(field);chapter.paragraph_format.page_break_before=True
    settings=d.settings.element;upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
    d.save(out)
    return out


def pdf_build(md):
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    style=ParagraphStyle('CN',fontName='STSong-Light',fontSize=11,leading=19,wordWrap='CJK',spaceAfter=7,firstLineIndent=22)
    heads={k:ParagraphStyle('H'+str(k),parent=style,fontSize={1:19,2:15,3:12}[k],leading=24,spaceBefore=12,firstLineIndent=0,keepWithNext=True) for k in (1,2,3)}
    caption=ParagraphStyle('caption',parent=style,fontSize=10,alignment=TA_CENTER,firstLineIndent=0)
    story=[];text=md.read_text(encoding='utf-8');lines=text.splitlines();buf=[]
    def flush():
        if buf:
            s=' '.join(buf);s=re.sub(r'`([^`]+)`',r'\1',s)
            story.append(Paragraph(escape(s),style));buf.clear()
    for line in lines:
        if not line.strip():flush();continue
        h=re.match(r'^(#{1,3})\s+(.*)',line)
        im=re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)',line)
        if h:
            flush();level=len(h[1]);title=h[2]
            if title in ['摘要','Abstract'] or title.startswith('1　研究'):story.append(PageBreak())
            story.append(Paragraph(escape(title),heads[level]))
        elif im:
            flush();p=(md.parent/im[2]).resolve();w,h=PILImage.open(p).size
            width=15.8*cm;height=width*h/w
            if height>20*cm:width*=20*cm/height;height=20*cm
            story.append(Image(str(p),width=width,height=height));story.append(Spacer(1,5))
        elif line.startswith('图 '):
            flush();story.append(Paragraph(escape(line),caption))
        else:buf.append(line)
    flush()
    out=md.with_suffix('.pdf')
    def page(canvas,doc):
        canvas.setFont('STSong-Light',9);canvas.drawCentredString(10.5*cm,1.3*cm,str(doc.page))
    SimpleDocTemplate(str(out),pagesize=(21*cm,29.7*cm),topMargin=2.2*cm,bottomMargin=2.2*cm,leftMargin=2.6*cm,rightMargin=2.6*cm,title=md.stem,author='课程设计材料').build(story,onFirstPage=page,onLaterPages=page)
    return out


def main():
    for name in ['课程设计报告.md','操作手册.md']:
        print(docx_build(ROOT/'docs'/name))
    print(pdf_build(ROOT/'docs'/'课程设计报告.md'))
    s=(ROOT/'docs'/'课程设计报告.md').read_text()
    b=s.split('## 1',1)[1].split('## 参考文献')[0]
    counts={'main_body_chinese_characters':len(re.findall('[\u4e00-\u9fff]',b)),
            'main_body_non_whitespace_characters':len(re.sub(r'\s','',b)),
            'method':'Chapter 1 through conclusions; excludes abstracts, references and appendices.'}
    import json
    (ROOT/'results'/'report_word_count.json').write_text(json.dumps(counts,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
